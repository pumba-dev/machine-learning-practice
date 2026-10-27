"""Monta slides/apresentacao.pptx a partir do template SBC e de slides/data/slides_data.json.

Passos, a partir da raiz do projeto:

    python scripts/slides_data.py       # números dos slides (seed 42 + 10 seeds)
    python scripts/build_slides.py      # slides/apresentacao.pptx e slides/notas.md
    powershell -File scripts/export_slides_pdf.ps1 slides/apresentacao.pptx   # PDF para validar

O template da SBC é o tema padrão do Office mais duas imagens de fundo; o fundo de conteúdo entra
em todos os slides. Gráficos e tabela são objetos nativos do PowerPoint, editáveis.
"""
import io
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slides_notes import build_notes

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "slides" / "template" / "Template Apresentações SBC.pptx"
DATA = ROOT / "slides" / "data" / "slides_data.json"
OUT = ROOT / "slides" / "apresentacao.pptx"
NOTES_MD = ROOT / "slides" / "notas.md"

# Divisão da fala em 3 blocos (tempo: 4:00, 3:00 e 3:00). Para trocar quem fala, mude só os nomes.
SPEAKERS = {1: "Thiago", 2: "Paulo", 3: "Bernan"}
BLOCK_OF_SLIDE = [1, 1, 1, 1, 2, 2, 2, 3, 3, 3]
BLOCK_TITLE = {1: "O problema, a base e os descritores",
               2: "Atributos do SMILES, avaliar sem se enganar e o resultado principal",
               3: "O que a árvore aprendeu, onde erra e a conclusão"}
# Frase de passagem, dita ao fim do último slide de um bloco; índice = posição do slide (0 a 9).
HANDOFF = {
    3: (2, "Esses cinco números são tudo que o modelo vê de cada molécula. Será que bastam?"),
    6: (3, "Vamos ver agora o que essa árvore aprendeu e onde ela erra."),
}

# Paleta: verdes do fundo da SBC, ouro de apoio e as cores de classe dos gráficos do relatório.
# Acessibilidade (projetor, sala iluminada, plateia com baixa visão): texto de no mínimo 18 pt, cores
# escuras sobre fundo claro e barras com contraste de pelo menos 3:1 contra o branco.
MIN_FONT = 18
INK = "1F2A2E"
GREEN = "09846C"
GREEN_DK = "0F4334"
TINT = "E3F0EC"
MUTED = "3C4649"
GRAY_BAR = "8A9795"
TEAL_BAR = "4E9C8A"
GOLD = "E3B505"
RED = "B23A2E"
RED_TINT = "F6E4E0"
CLASS_COLOR = {"EGFR": "2A78D6", "BACE1": "EB6834", "DRD2": "1BAF7A", "HDAC1": "EDA100"}
CLASS_TINT = {"EGFR": "D9E7F8", "BACE1": "FBE1D5", "DRD2": "D2F0E4", "HDAC1": "FBEFC9"}

FEATURE_LABEL = {"Molecular Weight": "Massa molecular", "LogP": "LogP", "HBA": "HBA", "HBD": "HBD", "TPSA": "TPSA"}
ATTRIBUTE_LABEL = {"n_aromatic": "átomos aromáticos", "TPSA": "TPSA", "HBD": "HBD", "n_rings": "anéis"}


def br(value, digits=1):
    """Número com vírgula decimal."""
    return f"{value:.{digits}f}".replace(".", ",")


def thousands(value):
    return f"{value:,}".replace(",", ".")


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


# --- Texto e formas -----------------------------------------------------------------------------

def write(frame, lines, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, font=None):
    """Escreve `lines` num text frame. Cada linha é str ou dict(text|runs, size, bold, color, align,
    space_after, bullet, font). Um run é str ou (str, dict(size, bold, color, font))."""
    first = True
    for line in lines:
        item = {"text": line} if isinstance(line, str) else line
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.alignment = item.get("align", align)
        if "space_after" in item:
            paragraph.space_after = Pt(item["space_after"])
        runs = item.get("runs") or [item.get("text", "")]
        for run_spec in runs:
            text, style = (run_spec, {}) if isinstance(run_spec, str) else run_spec
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(max(MIN_FONT, style.get("size", item.get("size", size))))
            run.font.language_id = MSO_LANGUAGE_ID.BRAZILIAN_PORTUGUESE
            run.font.bold = style.get("bold", item.get("bold", bold))
            run.font.color.rgb = rgb(style.get("color", item.get("color", color)))
            name = style.get("font", item.get("font", font))
            if name:
                run.font.name = name
        if item.get("bullet"):
            bullet(paragraph)


def bullet(paragraph, indent=0.25):
    """Marcador de lista de verdade (propriedade do parágrafo), não um caractere no texto."""
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", str(int(Inches(indent))))
    properties.set("indent", str(-int(Inches(indent))))
    marker = etree.SubElement(properties, qn("a:buChar"))
    marker.set("char", "•")


def text_box(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, name=None, **style):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    write(frame, lines, **style)
    if name:
        box.name = name
    return box


def shape(slide, kind, x, y, w, h, fill, line=None, radius=None, name=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.shadow.inherit = False
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid()
        item.fill.fore_color.rgb = rgb(fill)
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = rgb(line)
        item.line.width = Pt(1.25)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = radius
    if name:
        item.name = name
    return item


def card(slide, x, y, w, h, fill, lines=None, name=None, anchor=MSO_ANCHOR.MIDDLE, pad=0.18, **style):
    """Retângulo arredondado, com texto opcional dentro."""
    item = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.08, name=name)
    if lines:
        frame = item.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = Inches(pad)
        frame.margin_top = frame.margin_bottom = Inches(0.06)
        write(frame, lines, **style)
    return item


def line_between(slide, x1, y1, x2, y2, color, width=1.5, dash=False, name=None):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if dash:
        connector.line.dash_style = MSO_LINE.DASH
    if name:
        connector.name = name
    return connector


def number_badge(slide, x, y, diameter, label, fill=GREEN):
    badge = shape(slide, MSO_SHAPE.OVAL, x, y, diameter, diameter, fill, name=f"Número {label}")
    frame = badge.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(frame, [label], size=26, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)


def alt(item, text):
    """Texto alternativo (leitor de tela) do primeiro objeto de desenho de `item`."""
    item._element.xpath(".//p:cNvPr")[0].set("descr", text)


# --- Gráficos ----------------------------------------------------------------------------------

def style_chart(chart, size=18):
    chart.has_title = False
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(INK)


def label_points(series, values, size=20):
    """Rótulo de cada ponto como texto, com vírgula decimal e o mesmo arredondamento do restante do
    texto. O formato numérico do gráfico segue o idioma de quem abre o arquivo e mostraria ponto."""
    for index, value in enumerate(values):
        label = series.points[index].data_label
        frame = label.text_frame
        frame.text = f"{br(value)}%"
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = rgb(INK)
        label.position = XL_LABEL_POSITION.OUTSIDE_END


def column_chart(slide, x, y, w, h, categories, series, colors, description, vmax=100):
    """Colunas agrupadas, valores em %; rótulo de dado em cada coluna."""
    data = CategoryChartData(number_format='0.0"%"')
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    alt(frame, description)
    chart = frame.chart
    style_chart(chart)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.gap_width = 25
    plot.overlap = -5
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '0.0"%"'
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(20)
    labels.font.bold = True
    for item, color, (_, values) in zip(plot.series, colors, series):
        item.format.fill.solid()
        item.format.fill.fore_color.rgb = rgb(color)
        label_points(item, values, 18)
    axis = chart.value_axis
    axis.minimum_scale, axis.maximum_scale, axis.major_unit = 0, vmax, 20
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = rgb("DDE5E3")
    axis.format.line.fill.background()
    axis.tick_labels.number_format = '0"%"'
    axis.tick_labels.number_format_is_linked = False
    axis.tick_labels.font.size = Pt(MIN_FONT)
    axis.tick_labels.font.color.rgb = rgb(MUTED)
    chart.category_axis.tick_labels.font.size = Pt(MIN_FONT)
    chart.category_axis.format.line.color.rgb = rgb(GRAY_BAR)
    return chart


def bar_chart(slide, x, y, w, h, categories, values, colors, description, vmax=118):
    """Barras horizontais de uma série, primeira categoria no topo; sem eixo de valores (há rótulos).
    O máximo do eixo passa de 100 para sobrar espaço ao rótulo das barras mais longas."""
    data = CategoryChartData(number_format='0.0"%"')
    data.categories = categories
    data.add_series("Acurácia", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    alt(frame, description)
    chart = frame.chart
    style_chart(chart)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '0.0"%"'
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(20)
    labels.font.bold = True
    series = plot.series[0]
    for index, color in enumerate(colors):
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
    label_points(series, values)
    chart.category_axis.reverse_order = True
    chart.category_axis.tick_labels.font.size = Pt(MIN_FONT)
    chart.category_axis.format.line.color.rgb = rgb(GRAY_BAR)
    axis = chart.value_axis
    axis.minimum_scale, axis.maximum_scale = 0, vmax
    axis.has_major_gridlines = False
    axis.visible = False
    return chart


# --- Estrutura comum ---------------------------------------------------------------------------

SLIDE_NUMBER_XML = (
    '<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="{shape_id}" name="Número do slide"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="sldNum" sz="quarter" idx="12"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm></p:spPr>'
    '<p:txBody><a:bodyPr lIns="0" rIns="0" anchor="ctr"/><a:lstStyle/>'
    '<a:p><a:pPr algn="r"/>'
    '<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
    '<a:rPr lang="pt-BR" sz="1800" b="1"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:rPr>'
    '<a:t>‹#›</a:t></a:fld>'
    '<a:endParaRPr lang="pt-BR" sz="1800"/></a:p></p:txBody></p:sp>'
)


def add_slide_number(slide):
    """Número do slide como campo do PowerPoint (acompanha a ordem dos slides). Fica à esquerda do
    logo da SBC (x > 12,3) e acima da linha dourada do fundo (y ≈ 7,17)."""
    xml = SLIDE_NUMBER_XML.format(shape_id=slide.shapes._next_shape_id, x=int(Inches(10.9)), y=int(Inches(6.75)),
                                  w=int(Inches(1.2)), h=int(Inches(0.4)), color=MUTED)
    slide.shapes._spTree.append(parse_xml(xml))


def compose_notes(entry, speaker, handoff, closing, markdown=False):
    """Texto das notas de um slide: roteiro (fala), material de estudo e perguntas prováveis.

    O PowerPoint mostra as notas como texto simples, então o formato do .pptx usa rótulos em maiúsculas;
    o notas.md usa títulos e marcadores de Markdown."""
    spoken = f"[{speaker} · {entry['tempo']}] {entry['fala']}"
    extras = []
    if handoff:
        extras.append(f"Passagem para {handoff[0]}: “{handoff[1]}”")
    if closing:
        extras.append(closing)
    if markdown:
        parts = [f"**Roteiro (o que falar)** · {speaker} · {entry['tempo']}\n\n{entry['fala']}"]
        parts += [f"> {line}" for line in extras]
        parts.append("**Para estudar** (apoio; não faz parte da fala)\n\n" + "\n".join(f"- {line}" for line in entry["estudo"]))
        parts.append("**Perguntas prováveis**\n\n" + "\n\n".join(f"- **P:** {q}  \n  **R:** {a}" for q, a in entry["perguntas"]))
        return "\n\n".join(parts)
    parts = ["ROTEIRO (o que falar)\n" + spoken]
    parts += extras
    parts.append("────────────────────\nPARA ESTUDAR (apoio; não faz parte da fala)\n" + "\n".join(f"• {line}" for line in entry["estudo"]))
    parts.append("PERGUNTAS PROVÁVEIS\n" + "\n".join(f"P: {q}\nR: {a}\n" for q, a in entry["perguntas"]).rstrip())
    return "\n\n".join(parts)


class Deck:
    def __init__(self, data):
        self.script = build_notes(data)
        self.prs = Presentation(str(TEMPLATE))
        # O fundo da SBC vem como figura dos slides de exemplo; guarda o de conteúdo e esvazia o deck.
        self.art = self.prs.slides[1].shapes[0].image.blob
        slides = self.prs.slides._sldIdLst
        for slide_id in list(slides):
            self.prs.part.drop_rel(slide_id.rId)
            slides.remove(slide_id)
        self.layouts = {layout.name: layout for layout in self.prs.slide_layouts}
        self.notes = []

    def slide(self, layout, title, title_size=40):
        slide = self.prs.slides.add_slide(self.layouts[layout])
        background = slide.shapes.add_picture(io.BytesIO(self.art), 0, 0, self.prs.slide_width, self.prs.slide_height)
        background.name = "Fundo SBC"
        background._element.nvPicPr.cNvPr.set("descr", "Fundo decorativo do template da SBC")
        tree = slide.shapes._spTree
        tree.remove(background._element)
        tree.insert(2, background._element)  # atrás de tudo
        frame = slide.shapes.title.text_frame
        frame.text = title
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = 0  # alinha o título com o conteúdo, que começa em x = 0,92
        for run in frame.paragraphs[0].runs:
            run.font.size = Pt(title_size)
            run.font.bold = True  # a fonte leve do tema some no projetor; negrito pesa mais
            run.font.name = "Calibri"
            run.font.language_id = MSO_LANGUAGE_ID.BRAZILIAN_PORTUGUESE
            run.font.color.rgb = rgb(GREEN_DK)
        index = len(self.notes)
        block = BLOCK_OF_SLIDE[index]
        entry = self.script[index]
        handoff = None
        if index in HANDOFF:
            to_block, phrase = HANDOFF[index]
            handoff = (SPEAKERS[to_block], phrase)
        closing = f"{SPEAKERS[block]} conduz as perguntas." if index == len(BLOCK_OF_SLIDE) - 1 else None
        slide.notes_slide.notes_text_frame.text = compose_notes(entry, SPEAKERS[block], handoff, closing)
        self.notes.append((title, compose_notes(entry, SPEAKERS[block], handoff, closing, markdown=True)))
        if index > 0:  # o slide de título fica sem número
            add_slide_number(slide)
        return slide

    def save(self):
        OUT.parent.mkdir(parents=True, exist_ok=True)
        properties = self.prs.core_properties
        properties.title = "Classificação de alvos proteicos a partir de descritores moleculares do ChEMBL com árvores de decisão"
        properties.author = "Bernan R. do Nascimento; Paulo E. R. Araujo; Thiago R. P. de Carvalho"
        properties.language = "pt-BR"
        self.prs.save(str(OUT))
        lines = ["# Roteiro da apresentação (10 min)\n",
                 "Cada slide traz o roteiro (o que falar), um material de estudo e perguntas prováveis. O estudo e as "
                 "perguntas servem para dominar o assunto e não fazem parte da fala. As mesmas notas estão no .pptx.\n"]
        current = None
        for number, (title, notes) in enumerate(self.notes, 1):
            block = BLOCK_OF_SLIDE[number - 1]
            if block != current:
                current = block
                slides = [n for n, b in enumerate(BLOCK_OF_SLIDE, 1) if b == block]
                lines.append(f"# Bloco {block} · {SPEAKERS[block]} · slides {slides[0]} a {slides[-1]}: {BLOCK_TITLE[block]}\n")
            lines.append(f"## Slide {number}: {title}\n\n{notes}\n")
        NOTES_MD.write_text("\n".join(lines), encoding="utf-8")


# --- Slides ------------------------------------------------------------------------------------
# Tamanhos pensados para projetor: título 40 pt, texto de apoio de 18 a 20 pt, destaques de 24 a 66 pt.

def slide_title(deck, d):
    slide = deck.slide(
        "Slide de Título",
        "Classificação de alvos proteicos a partir de descritores moleculares do ChEMBL com árvores de decisão",
        title_size=40,
    )
    slide.shapes.title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    subtitle = slide.placeholders[1]
    frame = subtitle.text_frame
    frame.text = ""
    write(frame, [
        {"text": "Bernan R. do Nascimento · Paulo E. R. Araujo · Thiago R. P. de Carvalho", "size": 22, "bold": True, "space_after": 6},
        {"text": "Universidade Federal do Piauí (UFPI)", "size": 20},
        {"text": "PPGCC010 – Aprendizagem de Máquina (2026.2) · Prof. Vinicius Ponte Machado", "size": 20},
    ], align=PP_ALIGN.CENTER)


def slide_problem(deck, d):
    classes = d["classes"]
    maxima = max(d["descritor_sozinho"].values())
    macros = d["multiseed"]["macros"]
    slide = deck.slide("Somente Título", "Prever o alvo de uma molécula é mais difícil do que parece")
    text_box(slide, 0.92, 2.0, 5.4, 0.45, [f"4 alvos · {thousands(d['total_moleculas'])} moléculas"], size=24, bold=True, color=GREEN_DK)
    positions = [(0.92, 2.55), (3.7, 2.55), (0.92, 4.15), (3.7, 4.15)]
    for (name, info), (x, y) in zip(classes.items(), positions):
        card(slide, x, y, 2.6, 1.45, CLASS_TINT[name], [
            {"runs": [("●  ", {"color": CLASS_COLOR[name], "size": 24}), (name, {"bold": True, "size": 24})], "space_after": 0},
            {"text": f"{thousands(info['moleculas'])} moléculas", "size": 20},
            {"text": f"{br(info['percentual'])}%", "size": 20, "bold": True},
        ], name=f"Cartão {name}")
    card(slide, 0.92, 5.75, 5.38, 0.95, TINT,
         [f"Responder sempre EGFR acerta {br(d['piso'])}%: esse é o piso."], name="Piso", size=20, bold=True)

    text_box(slide, 6.9, 2.0, 5.1, 0.45, ["Três desafios"], size=24, bold=True, color=GREEN_DK)
    challenges = [
        ("Base de partida quebrada", "o arquivo do Kaggle não bate com a própria descrição"),
        ("Descritores fracos", f"um descritor sozinho acerta no máximo {br(maxima)}%"),
        ("Teste contaminado", f"{macros['gemeas']}% das moléculas de teste têm uma gêmea no treino"),
    ]
    for index, (head, detail) in enumerate(challenges):
        y = 2.6 + index * 1.4
        number_badge(slide, 6.9, y, 0.7, str(index + 1))
        text_box(slide, 7.8, y - 0.03, 4.2, 1.3, [
            {"text": head, "size": 24, "bold": True, "space_after": 2},
            {"text": detail, "size": 20},
        ], name=f"Desafio {index + 1}")


def slide_data(deck, d):
    slide = deck.slide("Somente Título", "A base do Kaggle estava quebrada, então reconstruímos do ChEMBL")
    card(slide, 0.92, 2.0, 3.9, 4.7, RED_TINT, None, name="Painel Kaggle")
    text_box(slide, 1.15, 2.12, 3.5, 0.45, ["Arquivo do Kaggle"], size=24, bold=True, color=RED)
    stats = [("1.333", "linhas (descritas: 8.711)"), ("0", "moléculas de EGFR"), ("95,5%", "das linhas são DRD2")]
    for index, (number, caption) in enumerate(stats):
        y = 2.75 + index * 1.3
        text_box(slide, 1.15, y, 3.5, 1.2, [
            {"text": number, "size": 44, "bold": True, "color": RED, "space_after": 0},
            {"text": caption, "size": 20},
        ], name=f"Kaggle {index + 1}")

    text_box(slide, 5.2, 2.0, 6.8, 0.45, ["Reconstrução do ChEMBL_37 (24/10/2026)"], size=24, bold=True, color=GREEN_DK)
    stages = [
        (58736, "medições de ligação"),
        (51359, "sem duplicata ou SMILES vazio"),
        (37424, "pares molécula-alvo (mediana)"),
        (27798, "pares com pChEMBL ≥ 6"),
        (27714, "moléculas de alvo único (−42)"),
        (27476, "com os 5 descritores (−238)"),
    ]
    for index, (count, caption) in enumerate(stages):
        y = 2.65 + index * 0.68
        width = max(1.6, 2.9 * count / stages[0][0])
        fill = GREEN_DK if index == len(stages) - 1 else GREEN
        bar = card(slide, 5.2, y, width, 0.56, fill, [thousands(count)], name=f"Etapa {index + 1}",
                   size=20, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER, pad=0.05)
        bar.adjustments[0] = 0.2
        text_box(slide, 8.3, y, 3.7, 0.56, [caption], anchor=MSO_ANCHOR.MIDDLE, size=20, name=f"Corte {index + 1}")


#: Descritor do ChEMBL, nome no slide, unidade, o que mede em poucas palavras e casas decimais.
DESCRIPTORS = [
    ("Molecular Weight", "Massa molecular", "dáltons", "Soma das massas dos átomos", 2),
    ("LogP", "LogP", "sem unidade", "Quanto prefere gordura a água", 2),
    ("HBA", "HBA", "átomos", "Aceitam ligação de hidrogênio", 0),
    ("HBD", "HBD", "grupos", "Doam ligação de hidrogênio", 0),
    ("TPSA", "TPSA", "Å²", "Área ocupada por átomos polares", 2),
]


def slide_descriptors(deck, d):
    example = d["exemplo"]
    values = example["descritores"]
    slide = deck.slide("Somente Título", "A base traz 5 descritores por molécula, calculados da estrutura")
    text_box(slide, 0.92, 2.0, 11.0, 0.45, [f"Exemplo: molécula {example['id']} ({example['classe']})"],
             size=24, bold=True, color=GREEN_DK)
    text_box(slide, 0.92, 2.5, 11.0, 0.4, ["Valores do ChEMBL; no próximo slide, a mesma molécula como texto"], size=20)
    for index, (key, label, unit, meaning, digits) in enumerate(DESCRIPTORS):
        y = 3.1 + index * 0.72
        card(slide, 0.92, y, 11.0, 0.64, TINT, None, name=f"Linha {label}")
        text_box(slide, 1.15, y, 2.6, 0.64, [label], anchor=MSO_ANCHOR.MIDDLE, size=22, bold=True, color=GREEN_DK, name=f"Nome {label}")
        text_box(slide, 3.8, y, 1.9, 0.64, [br(values[key], digits)], anchor=MSO_ANCHOR.MIDDLE, size=28, bold=True,
                 color=GREEN_DK, name=f"Valor {label}")
        text_box(slide, 5.7, y, 1.7, 0.64, [unit], anchor=MSO_ANCHOR.MIDDLE, size=20, name=f"Unidade {label}")
        text_box(slide, 7.5, y, 4.3, 0.64, [meaning], anchor=MSO_ANCHOR.MIDDLE, size=20, name=f"Significado {label}")


def slide_attributes(deck, d):
    single = d["descritor_sozinho"]
    ordered = sorted(single.items(), key=lambda item: item[1], reverse=True)
    example = d["exemplo"]
    counts = example["contagens"]
    slide = deck.slide("Somente Título", "Cinco descritores não bastam; contar o SMILES ajuda")
    text_box(slide, 0.92, 2.0, 5.3, 0.45, ["Acurácia com um só descritor"], size=24, bold=True, color=GREEN_DK)
    colors = [GREEN if value == ordered[0][1] else TEAL_BAR for _, value in ordered]
    bar_chart(slide, 0.92, 2.5, 5.3, 3.4, [FEATURE_LABEL[name] for name, _ in ordered],
              [value for _, value in ordered], colors,
              "Gráfico de barras: acurácia de uma árvore com um só descritor. HBD 45,1%, TPSA 44,9%, HBA 43,3%, "
              "LogP 38,7% e massa molecular 37,8%.", vmax=62)
    text_box(slide, 0.92, 6.0, 5.3, 0.45, [
        {"runs": [("Piso: ", {"bold": True}), (f"sempre EGFR = {br(d['piso'])}%", {})]},
    ], size=20, name="Piso")

    text_box(slide, 6.7, 2.0, 5.3, 0.45, ["12 atributos do SMILES"], size=24, bold=True, color=GREEN_DK)
    text_box(slide, 6.7, 2.5, 5.3, 0.4, ["Texto que descreve a molécula"], size=20)
    smiles = example["smiles"]
    cut = 24
    card(slide, 6.7, 3.0, 5.3, 0.95, TINT, [smiles[:cut], smiles[cut:]], name="SMILES de exemplo",
         size=18, font="Courier New", align=PP_ALIGN.CENTER, pad=0.1)
    tiles = [("smiles_len", "comprimento"), ("n_N", "átomos de N"), ("n_O", "átomos de O"),
             ("n_rings", "anéis"), ("n_aromatic", "aromáticos"), ("n_double", "ligações duplas")]
    for index, (key, label) in enumerate(tiles):
        x = 6.7 + (index % 3) * 1.8
        y = 4.15 + (index // 3) * 1.3
        card(slide, x, y, 1.7, 1.2, "FFFFFF", [
            {"text": str(counts[key]), "size": 32, "bold": True, "color": GREEN, "align": PP_ALIGN.CENTER, "space_after": 0},
            {"text": label, "size": 18, "align": PP_ALIGN.CENTER},
        ], name=f"Contagem {label}", pad=0.03).line.color.rgb = rgb(GRAY_BAR)


def twin_panel(slide, x0, title, grouped):
    """Painel esquemático: 12 moléculas, 4 pares de gêmeas, divididas em treino e teste."""
    y0 = 2.0
    panel = card(slide, x0, y0, 5.2, 3.3, "F5FAF8", None, name=f"Painel {title}")
    panel.line.color.rgb = rgb(GRAY_BAR)
    alt(panel, f"Esquema da {title.lower()}: 12 moléculas, em 4 pares de gêmeas, divididas em treino e teste. "
               + ("Cada par de gêmeas fica do mesmo lado." if grouped
                  else "Dois pares de gêmeas ficam divididos entre treino e teste, ligados por linhas tracejadas."))
    text_box(slide, x0 + 0.2, y0 + 0.08, 4.8, 0.45, [title], size=22, bold=True, color=GREEN_DK)
    zone_y = y0 + 0.55
    card(slide, x0 + 0.2, zone_y, 3.5, 1.5, TINT, None, name="Zona de treino")
    card(slide, x0 + 3.9, zone_y, 1.1, 1.5, "FFF1CC", None, name="Zona de teste")
    text_box(slide, x0 + 0.2, zone_y + 1.53, 3.5, 0.35, ["treino"], size=18, align=PP_ALIGN.CENTER)
    text_box(slide, x0 + 3.9, zone_y + 1.53, 1.1, 0.35, ["teste"], size=18, align=PP_ALIGN.CENTER)

    train = lambda col, row: (x0 + 0.4 + col * 0.85, zone_y + 0.22 + row * 0.75)  # noqa: E731
    test = lambda col, row: (x0 + 4.02 + col * 0.5, zone_y + 0.22 + row * 0.75)  # noqa: E731
    blue, orange, green, amber, gray = (CLASS_COLOR["EGFR"], CLASS_COLOR["BACE1"], CLASS_COLOR["DRD2"], CLASS_COLOR["HDAC1"], "7A8785")
    if grouped:
        molecules = [
            (train(0, 0), blue), (train(1, 0), blue), (train(0, 1), orange), (train(1, 1), orange),
            (train(2, 0), green), (train(3, 0), green), (train(2, 1), gray), (train(3, 1), gray),
            (test(0, 0), amber), (test(0, 1), amber), (test(1, 0), gray), (test(1, 1), gray),
        ]
        pairs = [(0, 1), (2, 3), (4, 5), (8, 9)]
    else:
        molecules = [
            (train(0, 0), blue), (train(1, 0), blue), (train(0, 1), orange), (train(1, 1), orange),
            (train(3, 0), green), (train(3, 1), amber), (train(2, 0), gray), (train(2, 1), gray),
            (test(0, 0), green), (test(0, 1), amber), (test(1, 0), gray), (test(1, 1), gray),
        ]
        pairs = [(0, 1), (2, 3), (4, 8), (5, 9)]
    diameter = 0.36
    for first, second in pairs:
        (xa, ya), (xb, yb) = molecules[first][0], molecules[second][0]
        crossing = (first < 8) != (second < 8)
        line_between(slide, xa + diameter / 2, ya + diameter / 2, xb + diameter / 2, yb + diameter / 2,
                     RED if crossing else MUTED, width=2.5 if crossing else 1.5, dash=crossing, name="Ligação entre gêmeas")
    for (x, y), color in molecules:
        shape(slide, MSO_SHAPE.OVAL, x, y, diameter, diameter, color, line="FFFFFF", name="Molécula")
    caption = ("Gêmeas sempre do mesmo lado: o teste exige generalizar" if grouped
               else "Gêmeas dos dois lados: acertar o teste vira lembrar o treino")
    text_box(slide, x0 + 0.2, y0 + 2.5, 4.8, 0.75, [caption], size=18, bold=True, color=GREEN_DK, name="Legenda do painel")


def slide_split(deck, d):
    macros = d["multiseed"]["macros"]
    slide = deck.slide("Somente Título", "A divisão aleatória deixa o teste ver a resposta")
    twin_panel(slide, 0.92, "Divisão aleatória", grouped=False)
    twin_panel(slide, 6.8, "Divisão agrupada", grouped=True)
    text_box(slide, 0.92, 5.4, 11.1, 0.4, ["Mesma cor: gêmeas (mesmos 5 descritores) · esquema ilustrativo"],
             size=18, name="Legenda do esquema")
    text_box(slide, 0.92, 5.9, 2.2, 0.85, [{"text": f"{macros['gemeas']}%", "size": 44, "bold": True, "color": GREEN}],
             anchor=MSO_ANCHOR.MIDDLE, name="Destaque gêmeas")
    text_box(slide, 3.15, 5.9, 3.0, 0.85, ["do teste tem gêmea no treino (aleatória)"], anchor=MSO_ANCHOR.MIDDLE,
             size=18, name="Legenda do destaque")
    card(slide, 6.8, 5.9, 5.2, 0.85, TINT, ["3 conjuntos × 2 divisões = 6 cenários", "repetidos com 10 seeds"],
         name="Cenários", size=20, bold=True, color=GREEN_DK, align=PP_ALIGN.CENTER)


def slide_results(deck, d):
    multi = d["multiseed"]
    macros = multi["macros"]
    scenarios = multi["cenarios"]
    keys = ["descritores", "só SMILES", "descritores + SMILES"]
    slide = deck.slide("Somente Título", f"Contar o SMILES eleva a acurácia em {macros['ganhoAmbos']}%")
    text_box(slide, 0.92, 2.0, 6.1, 0.45, ["Acurácia por conjunto de atributos"], size=24, bold=True, color=GREEN_DK)
    column_chart(
        slide, 0.92, 2.5, 6.5, 4.2,
        ["Descritores (5)", "SMILES (12)", "Ambos (17)"],
        [("Divisão aleatória", [scenarios[k]["aleatoria"] for k in keys]),
         ("Divisão agrupada", [scenarios[k]["agrupada"] for k in keys])],
        [GRAY_BAR, GREEN],
        "Gráfico de colunas: acurácia no teste, média de 10 seeds. Descritores: aleatória "
        f"{macros['accAleatDesc']}%, agrupada {macros['accGrupDesc']}%. SMILES: {macros['accAleatSmiles']}% e "
        f"{macros['accGrupSmiles']}%. Ambos: {macros['accAleatAmbos']}% e {macros['accGrupAmbos']}%.",
    )
    text_box(slide, 7.6, 2.0, 4.4, 1.1, [{"text": f"+{macros['ganhoAmbos']}%", "size": 66, "bold": True, "color": GREEN}], name="Ganho")
    text_box(slide, 7.6, 3.15, 4.4, 0.8, [f"ganho relativo com as contagens de SMILES ({macros['accGrupDesc']}% → {macros['accGrupAmbos']}%)"],
             size=20, name="Legenda do ganho")
    text_box(slide, 7.6, 4.1, 4.4, 0.4, ["Superestimativa da divisão aleatória"], size=20, bold=True, color=GREEN_DK)
    for index, (value, label) in enumerate([(macros["gapDesc"], "descritores"), (macros["gapSmiles"], "SMILES"), (macros["gapAmbos"], "ambos")]):
        text_box(slide, 7.6, 4.6 + index * 0.55, 4.4, 0.5, [
            {"runs": [(f"{value}%", {"size": 28, "bold": True, "color": GREEN}), (f"  {label}", {"size": 20})]},
        ], anchor=MSO_ANCHOR.MIDDLE, name=f"Superestimativa {label}")
    text_box(slide, 7.6, 6.3, 4.4, 0.45, [f"Média de 10 seeds; desvio até {br(multi['dp_maximo'])}%"], size=18, name="Nota")


def slide_reference(deck, d):
    multi = d["multiseed"]
    macros = multi["macros"]
    reference = multi["referencia"]
    tree = d["arvore"]
    left, right = tree["esquerda"], tree["direita"]
    root_threshold = br(tree["raiz"]["limiar"])
    slide = deck.slide("Somente Título", f"A Random Forest supera a árvore em {macros['ganhoRF']}%, mas só a árvore tem regras legíveis")
    text_box(slide, 0.92, 2.0, 5.3, 0.45, ["Acurácia na divisão agrupada"], size=24, bold=True, color=GREEN_DK)
    names = ["Sempre EGFR", "Árvore de decisão", "Gradient boosting", "Random Forest"]
    values = [reference["Dummy (classe majoritária)"], reference["Árvore de decisão"],
              reference["HistGradientBoosting"], reference["Random Forest (300 árvores)"]]
    bar_chart(slide, 0.92, 2.55, 5.3, 3.9, names, values, [GRAY_BAR, GREEN, TEAL_BAR, TEAL_BAR],
              f"Gráfico de barras: acurácia na divisão agrupada. Sempre EGFR {br(values[0])}%, árvore de decisão "
              f"{br(values[1])}%, gradient boosting {br(values[2])}%, Random Forest {br(values[3])}%.")

    text_box(slide, 6.6, 2.0, 5.4, 0.45, ["Primeiros níveis da árvore"], size=24, bold=True, color=GREEN_DK)
    card(slide, 7.8, 2.6, 3.0, 0.8, GREEN, [f"TPSA ≤ {root_threshold}?"], name="Raiz",
         size=24, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    line_between(slide, 9.3, 3.4, 7.9, 4.1, MUTED, width=2, name="Ramo sim")
    line_between(slide, 9.3, 3.4, 10.7, 4.1, MUTED, width=2, name="Ramo não")

    def question(node):
        pergunta = node["pergunta"]
        # Espaços sem quebra: a pergunta inteira cai junta na linha de baixo.
        return f"depois: aromáticos ≤ {int(pergunta['limiar'])}?"

    card(slide, 6.6, 4.1, 2.6, 1.75, CLASS_TINT["DRD2"], [
        {"text": f"Sim: {br(left['fatia'] * 100)}%", "size": 20, "bold": True, "space_after": 2},
        {"text": f"{br(left['fatia_classe'] * 100, 0)}% são {left['classe']}", "size": 20, "space_after": 2},
        {"text": question(left), "size": 18},
    ], name="Ramo sim", anchor=MSO_ANCHOR.TOP, pad=0.12)
    card(slide, 9.4, 4.1, 2.6, 1.75, "FFFFFF", [
        {"text": f"Não: {br(right['fatia'] * 100)}%", "size": 20, "bold": True, "space_after": 2},
        {"text": f"{br(right['fatia_classe'] * 100, 0)}% {right['classe']} (mistura)", "size": 20, "space_after": 2},
        {"text": question(right), "size": 18},
    ], name="Ramo não", anchor=MSO_ANCHOR.TOP, pad=0.12).line.color.rgb = rgb(GRAY_BAR)
    text_box(slide, 6.6, 6.0, 5.4, 0.75, [
        f"Árvore completa: {thousands(tree['folhas'])} folhas. TPSA e aromáticos lideram a importância.",
    ], size=18, name="Legenda da árvore")


def slide_errors(deck, d):
    matrix = d["matriz_confusao"]
    classes = matrix["classes"]
    values = matrix["valores"]
    errors = d["erros"]
    slide = deck.slide("Somente Título", "A árvore erra entre HDAC1 e EGFR e não sabe que errou")
    text_box(slide, 0.92, 2.0, 5.8, 0.45, ["Matriz de confusão no teste (seed 42)"], size=24, bold=True, color=GREEN_DK)
    table_shape = slide.shapes.add_table(5, 5, Inches(0.92), Inches(2.55), Inches(5.8), Inches(3.5))
    table_shape.name = "Matriz de confusão"
    alt(table_shape, "Matriz de confusão no teste, em contagens. Linhas: classe real; colunas: classe prevista. "
                     + " ".join(f"{classes[i]}: " + ", ".join(f"{classes[j]} {values[i][j]}" for j in range(4)) + "." for i in range(4)))
    table = table_shape.table
    table.first_row = False
    table.horz_banding = False
    table.columns[0].width = Inches(1.4)
    for column in range(1, 5):
        table.columns[column].width = Inches(1.1)
    for row in table.rows:
        row.height = Inches(0.7)

    def fill_cell(row, column, text, fill, color=INK, bold=False):
        cell = table.cell(row, column)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.04)
        frame = cell.text_frame
        frame.text = ""
        write(frame, [text], size=20, bold=bold, color=color, align=PP_ALIGN.CENTER)

    fill_cell(0, 0, "", "FFFFFF")
    for index, name in enumerate(classes):
        fill_cell(0, index + 1, name, "FFFFFF", bold=True)
        fill_cell(index + 1, 0, name, "FFFFFF", bold=True)
        total = sum(values[index])
        for column, count in enumerate(values[index]):
            share = count / total
            # Branco → verde pela proporção da linha; texto branco quando o fundo escurece.
            mix = [round(255 + (c - 255) * share) for c in (0x09, 0x84, 0x6C)]
            fill_cell(index + 1, column + 1, thousands(count), "%02X%02X%02X" % tuple(mix),
                      color="FFFFFF" if share > 0.5 else INK, bold=share > 0.5)
    text_box(slide, 0.92, 6.15, 5.8, 0.4, ["Linha: classe real · coluna: prevista"], size=18, name="Legenda da matriz")

    text_box(slide, 7.1, 2.0, 4.9, 1.3, [
        {"text": f"{br(errors['par_hdac1_egfr'])}%", "size": 48, "bold": True, "color": GREEN, "space_after": 0},
        {"text": "dos erros são HDAC1 ↔ EGFR", "size": 22},
    ], name="Destaque do par")
    text_box(slide, 7.1, 3.4, 4.9, 1.7, [
        {"text": "Por que erra?", "size": 22, "bold": True, "space_after": 4},
        {"text": "Os atributos tirados do SMILES são só contagens e podem não bastar para representar a molécula", "size": 20},
    ], name="Limite dos atributos")
    text_box(slide, 7.1, 5.2, 4.9, 1.5, [
        {"text": f"{br(errors['previsoes_com_confianca_100'], 0)}%", "size": 48, "bold": True, "color": GREEN, "space_after": 0},
        {"text": f"de confiança nas {thousands(errors['teste'])} previsões; {br(errors['percentual_erro'])}% erradas", "size": 20},
    ], name="Destaque da confiança")


def slide_conclusion(deck, d):
    macros = d["multiseed"]["macros"]
    slide = deck.slide("Somente Título", "Conclusões e limites")
    cards = [
        (f"{macros['accGrupAmbos']}%", "acurácia da árvore na divisão agrupada, com 17 atributos"),
        (f"+{macros['ganhoAmbos']}%", "ganho relativo ao contar o SMILES, sobre só os descritores"),
        (f"{macros['gapMin']}% a {macros['gapMax']}%", "quanto a divisão aleatória superestima a acurácia"),
    ]
    for index, (number, caption) in enumerate(cards):
        card(slide, 0.92 + index * 3.78, 2.0, 3.55, 2.2, TINT, [
            {"text": number, "size": 40, "bold": True, "color": GREEN_DK, "align": PP_ALIGN.CENTER, "space_after": 4},
            {"text": caption, "size": 18, "align": PP_ALIGN.CENTER},
        ], name=f"Número {index + 1}")
    text_box(slide, 0.92, 4.45, 11.0, 1.9, [
        {"text": "Limites", "size": 24, "bold": True, "color": GREEN_DK, "space_after": 4},
        {"text": "Agrupa só descritores idênticos, não famílias de moléculas parecidas", "size": 22, "bullet": True, "space_after": 4},
        {"text": "Escolhe sempre uma de 4 classes", "size": 22, "bullet": True, "space_after": 4},
        {"text": "Sem validação externa", "size": 22, "bullet": True},
    ], name="Limites")
    text_box(slide, 0.92, 6.45, 9.5, 0.4, [
        {"runs": [("Código e dados: ", {"bold": True}), ("github.com/pumba-dev/machine-learning-practice", {}), ("  ·  Perguntas?", {"bold": True, "color": GREEN})]},
    ], size=20, name="Repositório")


def audit_fonts(minimum=MIN_FONT):
    """Confere que nenhum texto, tabela ou gráfico do arquivo gravado fica abaixo do tamanho mínimo."""
    prs = Presentation(str(OUT))
    found = []
    for number, slide in enumerate(prs.slides, 1):
        for item in slide.shapes:
            frames = []
            if item.has_text_frame:
                frames.append(item.text_frame)
            if getattr(item, "has_table", False) and item.has_table:
                frames += [cell.text_frame for row in item.table.rows for cell in row.cells]
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None and run.font.size.pt < minimum:
                            found.append((number, item.name, run.font.size.pt, run.text[:30]))
            if getattr(item, "has_chart", False) and item.has_chart:
                for size in item.chart._chartSpace.xpath(".//@sz"):
                    if int(size) < minimum * 100:
                        found.append((number, item.name, int(size) / 100, "gráfico"))
    return found


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    deck = Deck(data)
    for build in (slide_title, slide_problem, slide_data, slide_descriptors, slide_attributes, slide_split,
                  slide_results, slide_reference, slide_errors, slide_conclusion):
        build(deck, data)
    deck.save()
    print("gravado:", OUT)
    small = audit_fonts()
    print("texto abaixo de", MIN_FONT, "pt:", small if small else "nenhum")


if __name__ == "__main__":
    main()
